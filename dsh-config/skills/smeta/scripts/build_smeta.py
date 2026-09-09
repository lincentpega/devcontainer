#!/usr/bin/env python3
"""Build a restaurant смета (bill breakdown) as an Excel matrix.

Matches the format of assets/examples/*.xlsx:
  row 1        header  Позиция | Сумма | <person...>
  rows 2..N    item    name | =SUM(...) | per-person amounts (0 if none)
  blank row
  Полная сумма        per person =SUM(col);  Сумма = SUM(B items)
  Обслуживание N%     per person = CEILING(full*N/100,1) or exact
  ИТОГО               per person = full + service;  Сумма = B full + B service

Adapt `people` and `items`, set `OUT`, then run:
    python3 build_smeta.py
Requires openpyxl.
"""
import math
import openpyxl
from openpyxl.styles import Font, Alignment

# ----------------------------------------------------------------------------
# DATA — the only thing to edit per job.
# ----------------------------------------------------------------------------
OUT = "smeta.xlsx"
SERVICE_PERCENT = 12          # "Обслуживание N%"
ROUND_UP = True               # True -> CEILING service per person; False -> exact 12%
people = ["Я", "Алексей", "Диёр", "Баке", "Наиль", "Владислав", "Ирсен"]

# items: (receipt line, {person: amount}). Amount 0 is fine (write it or omit).
items = [
    ("Туй оши порция",      {"Баке": 50000, "Владислав": 50000, "Ирсен": 50000}),
    ("доп.гушт (туй оши)",  {"Владислав": 30000}),
    ("Чойхона ош порция",   {"Диёр": 50000}),
    ("Бедана тухум (2 шт)", {"Ирсен": 4000}),
    ("Зигир ош порция",     {"Я": 52000}),
    ("доп.гушт (зигир ош)", {"Я": 30000}),
    ("Катикли салат",       {"Ирсен": 12000}),
    ("Ачичук",              {"Я": 12000, "Алексей": 12000, "Наиль": 12000}),
    ("Туй оши 0,7 порция",  {"Алексей": 48000}),
    ("Особый ош 0,7 порция",{"Наиль": 48000}),
    ("доп.гушт (особый ош)",{"Наиль": 30000}),
    ("Нон бутун",           {"Диёр": 4000, "Ирсен": 4000}),
    ("Кук чой",             {"Я": 1667, "Алексей": 1667, "Диёр": 1667, "Баке": 1667,
                             "Владислав": 1667, "Ирсен": 1665}),
]
# ----------------------------------------------------------------------------


def build(out, service_percent, round_up):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Лист1"
    first = 3                     # first person column = C
    last = first + len(people) - 1
    sum_col = 2                   # B
    bold = Font(name="Arial", bold=True)
    norm = Font(name="Arial")
    h_align = Alignment(horizontal="left", vertical="center")
    v_align = Alignment(vertical="center")

    def L(c):
        return openpyxl.utils.get_column_letter(c)

    ws.cell(1, 1, "Позиция")
    ws.cell(1, sum_col, "Сумма")
    for j, p in enumerate(people):
        ws.cell(1, first + j, p)
    for c in range(1, last + 1):
        cell = ws.cell(1, c)
        cell.font = bold
        cell.alignment = h_align

    row = 2
    for name, assigns in items:
        ws.cell(row, 1, name).font = norm
        ws.cell(row, 1).alignment = v_align
        ws.cell(row, sum_col, f"=SUM({L(first)}{row}:{L(last)}{row})").font = norm
        ws.cell(row, sum_col).alignment = v_align
        for j, p in enumerate(people):
            cell = ws.cell(row, first + j, assigns.get(p, 0))
            cell.font = norm
            cell.alignment = v_align
        row += 1

    last_item = row - 1
    blank = row
    full = blank + 1
    service = full + 1
    total = service + 1

    ws.cell(full, 1, "Полная сумма").font = bold
    ws.cell(full, sum_col, f"=SUM(B2:B{last_item})").font = norm
    ws.cell(service, 1, f"Обслуживание {service_percent}%").font = bold
    ws.cell(service, sum_col, f"=SUM({L(first)}{service}:{L(last)}{service})").font = norm
    ws.cell(total, 1, "ИТОГО").font = bold
    ws.cell(total, sum_col, f"=B{full}+B{service}").font = norm

    for c in range(first, last + 1):
        col = L(c)
        ws.cell(full, c, f"=SUM({col}2:{col}{last_item})").font = norm
        if round_up:
            ws.cell(service, c, f"=CEILING({col}{full}*{service_percent}/100,1)").font = norm
        else:
            ws.cell(service, c, f"={col}{full}*{service_percent}/100").font = norm
        ws.cell(total, c, f"={col}{full}+{col}{service}").font = norm

    for k, v in {"A": 30, "B": 16, **{L(c): 12 for c in range(first, last + 1)}}.items():
        ws.column_dimensions[k].width = v
    for r in range(1, total + 1):
        ws.row_dimensions[r].height = 22.5
    ws.freeze_panes = "A2"

    wb.save(out)
    return last_item, full, service, total


def preview():
    sums = {p: 0 for p in people}
    for _, assigns in items:
        for p, a in assigns.items():
            sums[p] += a
    tf = ts = tt = 0
    rows = []
    for p in people:
        s = sums[p]
        sv = math.ceil(s * SERVICE_PERCENT / 100) if ROUND_UP else s * SERVICE_PERCENT / 100
        t = s + sv
        tf += s
        ts += sv
        tt += t
        rows.append((p, s, sv, t))
    print(f"{'person':>9} {'subtotal':>10} {'service':>10} {'ITOGO':>12}")
    for r in rows:
        print(f"{r[0]:>9} {r[1]:>10,} {r[2]:>10,.2f} {r[3]:>12,.2f}")
    print(f"{'TOTAL':>9} {tf:>10,} {ts:>10,.2f} {tt:>12,.2f}")


if __name__ == "__main__":
    last, full, service, total = build(OUT, SERVICE_PERCENT, ROUND_UP)
    print(f"wrote {OUT}  (items rows 2..{last}, subtotal {full}, service {service}, total {total})")
    preview()
